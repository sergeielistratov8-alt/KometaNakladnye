from pathlib import Path
import os
import openpyxl

# Create sample CSV
csv_path = Path('test_input.csv')
xlsx_path = Path('test_output.xlsx')

csv_content = '''
Артикул;Товар;Кол-во;Цена;Сумма
12345;Товар A; 10 ;1,234.56;12345.60
ABC-001;Товар B;;2,000; 
;Товар C;5; 300,5;1502,5
'''.lstrip()

with open(csv_path, 'w', encoding='utf-8') as f:
    f.write(csv_content)

import pandas as pd

df = pd.read_csv(csv_path, sep=';')

def convert_numeric_columns(df, columns=None):
    if columns is None:
        columns = ['Кол-во', 'Цена', 'Сумма', 'Артикул']
    cols = [c for c in columns if c in df.columns]
    for col in cols:
        try:
            s = df[col].astype(str).str.strip().replace('', pd.NA)
            s = s.str.replace(' ', '')
            def _normalize(x):
                if x is None:
                    return x
                xs = str(x)
                if ',' in xs and '.' in xs:
                    return xs.replace(',', '')
                if ',' in xs:
                    return xs.replace(',', '.')
                return xs
            s = s.apply(_normalize)
            conv = pd.to_numeric(s, errors='coerce')
            if conv.notna().any():
                mask = conv.notna()
                df[col] = df[col].astype(object)
                df.loc[mask, col] = conv[mask]
        except Exception:
            pass
    return df

df = convert_numeric_columns(df)

# Save to Excel using pandas (engine openpyxl)
df.to_excel(xlsx_path, index=False)

# Inspect output xlsx
wb = openpyxl.load_workbook(xlsx_path)
ws = wb.active

max_row = ws.max_row
max_col = ws.max_column

print('Worksheet size:', max_row, max_col)

for r in range(1, max_row+1):
    row_vals = []
    for c in range(1, max_col+1):
        cell = ws.cell(row=r, column=c)
        val = cell.value
        is_num = isinstance(val, (int, float))
        row_vals.append(f"{val!r}({'num' if is_num else 'str'})")
    print(' | '.join(row_vals))

# Clean up files (optional)
# os.remove(csv_path)
# os.remove(xlsx_path)
